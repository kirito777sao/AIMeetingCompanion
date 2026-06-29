"""
tests/test_report_export.py
-----------------------------
Phase 3 Bước 3/4 (F11) — test module `app/report_export.py`.

KHÔNG cần PySide6 (module export hoàn toàn độc lập Qt). CẦN `python-docx` +
`reportlab` cài thật (đã có trong requirements.txt) -- khác `whisper`/`groq`,
2 thư viện này chạy được HOÀN TOÀN OFFLINE (build file local, không gọi
network) nên test ở đây verify được THẬT, không cần mock/stub như
groq_provider.py.

⚠️ Phần PDF: sandbox này KHÔNG có font hệ thống Windows (`C:\\Windows\\
Fonts\\segoeui.ttf`) -- test "thiếu font" verify ĐÚNG hành vi raise lỗi rõ
ràng (hành vi THẬT sẽ xảy ra y như vậy nếu app chạy nhầm môi trường không có
font này). Để verify pipeline register-font + generate PDF chạy được VỀ MẶT
KỸ THUẬT, test "thành công" tạm trỏ `WINDOWS_FONT_REGULAR`/`_BOLD` sang 1
font TrueType khác có sẵn trong sandbox (DejaVu Sans, cũng hỗ trợ Unicode
đầy đủ dấu tiếng Việt) -- KHÔNG verify được đúng "Segoe UI" thật trên Windows
có hiển thị đúng dấu hay không, cần bạn tự mở file .pdf xuất ra THẬT bằng
Acrobat/trình đọc PDF trên Windows để xác nhận lại bằng mắt.

Chạy 2 test PDF THEO ĐÚNG THỨ TỰ dưới (test "thiếu font" TRƯỚC test "thành
công") -- reportlab cache font đã register theo TÊN trong toàn tiến trình
(`_register_pdf_fonts()` idempotent CÓ Ý, xem docstring), nên nếu chạy
ngược thứ tự, lần "thiếu font" sau sẽ KHÔNG raise nữa vì tên font đã được
register từ lần thành công trước đó trong CÙNG tiến trình Python.
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app.report_export as report_export  # noqa: E402
from app.report_export import ReportExportError, export_to_docx, export_to_pdf  # noqa: E402
from app.storage.db import ActionItemRecord, ReportRecord  # noqa: E402

_DEJAVU_REGULAR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
_DEJAVU_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def _sample_report(with_content: bool = True) -> ReportRecord:
    if not with_content:
        return ReportRecord(
            meeting_id=1, generated_at="2026-06-27T00:00:00+00:00",
            overview="", decisions=[], action_items=[],
        )
    return ReportRecord(
        meeting_id=1, generated_at="2026-06-27T00:00:00+00:00",
        overview="Đã thảo luận về deadline dự án & ngân sách Q3.",
        decisions=["Chốt deadline thứ Sáu tới", "Tăng ngân sách <quảng cáo> 10%"],
        action_items=[
            ActionItemRecord(text="Gửi báo giá cho khách hàng", done=True),
            ActionItemRecord(text="Họp lại với team vào tuần sau", done=False),
        ],
    )


def test_export_docx_co_du_noi_dung():
    report = _sample_report()
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "bao_cao.docx")
        export_to_docx(report, meeting_id=1, native_lang_label="Tiếng Việt", output_path=path)
        assert os.path.exists(path) and os.path.getsize(path) > 0

        import docx
        document = docx.Document(path)
        full_text = "\n".join(p.text for p in document.paragraphs)
        assert "deadline dự án" in full_text
        assert "Chốt deadline thứ Sáu tới" in full_text
        assert "[x]" in full_text and "Gửi báo giá" in full_text
        assert "[ ]" in full_text and "Họp lại với team" in full_text
        assert document.styles["Normal"].font.name == "Segoe UI"
    print("OK: test_export_docx_co_du_noi_dung")


def test_export_docx_rong_hien_khong_co():
    report = _sample_report(with_content=False)
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "bao_cao_rong.docx")
        export_to_docx(report, meeting_id=2, native_lang_label="Tiếng Việt", output_path=path)

        import docx
        document = docx.Document(path)
        full_text = "\n".join(p.text for p in document.paragraphs)
        assert full_text.count("(Không có)") == 3  # tóm tắt rỗng + decisions rỗng + action_items rỗng
    print("OK: test_export_docx_rong_hien_khong_co")


def test_export_docx_duong_dan_loi_raise_report_export_error():
    report = _sample_report()
    try:
        export_to_docx(report, meeting_id=1, native_lang_label="Tiếng Việt",
                        output_path="/duong-dan-khong-ton-tai/bao_cao.docx")
        raise AssertionError("Phải raise ReportExportError khi đường dẫn ghi không hợp lệ.")
    except ReportExportError:
        pass
    print("OK: test_export_docx_duong_dan_loi_raise_report_export_error")


def test_export_pdf_thieu_font_he_thong_raise_loi_ro_rang():
    # PHẢI chạy TRƯỚC test "thành công" dưới (xem docstring đầu file) --
    # sandbox này thật sự không có C:\Windows\Fonts\segoeui.ttf, đúng đường
    # dẫn mặc định của report_export.py, không cần monkeypatch gì ở đây.
    assert not os.path.exists(report_export.WINDOWS_FONT_REGULAR), (
        "Test này giả định sandbox KHÔNG có font Windows -- nếu thấy fail, "
        "môi trường chạy test đã đổi, cần xem lại."
    )
    report = _sample_report()
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "bao_cao.pdf")
        try:
            export_to_pdf(report, meeting_id=1, native_lang_label="Tiếng Việt", output_path=path)
            raise AssertionError("Phải raise ReportExportError khi thiếu font hệ thống.")
        except ReportExportError as exc:
            assert "font" in str(exc).lower()
        assert not os.path.exists(path), "KHÔNG được tạo ra file PDF lỗi câm khi thiếu font."
    print("OK: test_export_pdf_thieu_font_he_thong_raise_loi_ro_rang")


def test_export_pdf_voi_font_gia_lap_chay_duoc_thuc_te():
    # Monkeypatch sang DejaVu Sans (có sẵn trong sandbox) để verify pipeline
    # register-font + build PDF chạy được THẬT về mặt kỹ thuật -- xem giới
    # hạn đã ghi rõ ở docstring đầu file (KHÔNG verify được Segoe UI thật).
    original_regular = report_export.WINDOWS_FONT_REGULAR
    original_bold = report_export.WINDOWS_FONT_BOLD
    assert os.path.exists(_DEJAVU_REGULAR), "Cần fonts-dejavu-core trong sandbox để chạy test này."
    report_export.WINDOWS_FONT_REGULAR = _DEJAVU_REGULAR
    report_export.WINDOWS_FONT_BOLD = _DEJAVU_BOLD
    try:
        report = _sample_report()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "bao_cao.pdf")
            export_to_pdf(report, meeting_id=1, native_lang_label="Tiếng Việt", output_path=path)
            assert os.path.exists(path) and os.path.getsize(path) > 0
            with open(path, "rb") as f:
                header = f.read(5)
            assert header == b"%PDF-", f"File xuất ra phải là PDF hợp lệ, header thật: {header!r}"
    finally:
        report_export.WINDOWS_FONT_REGULAR = original_regular
        report_export.WINDOWS_FONT_BOLD = original_bold
    print("OK: test_export_pdf_voi_font_gia_lap_chay_duoc_thuc_te")


def test_export_pdf_ky_tu_dac_biet_khong_lam_loi_parse():
    # "&"/"<"/">" trong overview/decisions (đã có trong _sample_report()) --
    # nếu thiếu _escape() trong report_export.py, dòng build_PDF dưới sẽ
    # raise lỗi parse markup nội bộ của reportlab.
    # Font "SegoeUI" đã được register (idempotent theo TÊN) từ test thành
    # công ở trên, dù WINDOWS_FONT_REGULAR đã được restore lại path Windows
    # gốc (không tồn tại trong sandbox) -- export_to_pdf() dưới đây vẫn
    # chạy được vì _register_pdf_fonts() skip ngay khi thấy tên đã có trong
    # registry, không đọc lại path. Đây CHÍNH LÀ hành vi cố ý cần verify.
    report = _sample_report()
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "bao_cao_dac_biet.pdf")
        export_to_pdf(report, meeting_id=1, native_lang_label="Tiếng Việt", output_path=path)
        assert os.path.exists(path) and os.path.getsize(path) > 0
    print("OK: test_export_pdf_ky_tu_dac_biet_khong_lam_loi_parse")


if __name__ == "__main__":
    test_export_docx_co_du_noi_dung()
    test_export_docx_rong_hien_khong_co()
    test_export_docx_duong_dan_loi_raise_report_export_error()
    test_export_pdf_thieu_font_he_thong_raise_loi_ro_rang()
    test_export_pdf_voi_font_gia_lap_chay_duoc_thuc_te()
    test_export_pdf_ky_tu_dac_biet_khong_lam_loi_parse()
    print("\n=== TẤT CẢ TEST REPORT_EXPORT PASS ===")
